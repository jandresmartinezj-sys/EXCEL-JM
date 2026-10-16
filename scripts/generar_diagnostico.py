"""Genera la plantilla 'Diagnóstico financiero pyme' + propuesta comercial (1 página cada una)."""
import datetime as dt
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule, Rule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.styles.numbers import NumberFormat
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

AZUL = PatternFill("solid", fgColor="1F3864")
GRIS = PatternFill("solid", fgColor="F2F2F2")
INPUT = Font(color="0000FF")
BL = Font(color="FFFFFF", bold=True)
NEG = Font(bold=True)
B = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
COP = '#,##0;[Red]-#,##0'
WRAP = Alignment(wrap_text=True, vertical="top")


def head(ws, row, vals):
    for j, v in enumerate(vals, 1):
        c = ws.cell(row, j, v); c.fill, c.font = AZUL, BL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


wb = Workbook()

# ---------- Instrucciones ----------
ws = wb.active; ws.title = "Instrucciones"
for i, (t, b) in enumerate([
    ("DIAGNÓSTICO FINANCIERO PYME", True), ("", False),
    ("1. Llena la hoja 'Datos' (celdas en azul) con los estados financieros del último año o de los últimos 12 meses.", False),
    ("2. Responde el chequeo tributario (Sí / No / N/A).", False),
    ("3. La hoja 'Diagnostico' se genera sola: indicadores, semáforo, puntaje y acciones recomendadas. Está lista para imprimir o guardar como PDF en 1 página.", False),
    ("4. En 'Parametros' puedes ajustar los rangos verde/rojo según el sector del cliente.", False), ("", False),
    ("IMPORTANTE: los rangos son referencias generales, no estándares sectoriales. Ajústalos al sector y al tamaño del cliente.", True),
    ("Herramienta de apoyo; no reemplaza el juicio profesional ni una revisión completa de la información.", False),
], 1):
    ws.cell(i, 1, t).font = Font(bold=b, size=14 if i == 1 else 11)
ws.column_dimensions["A"].width = 120

# ---------- Datos ----------
d = wb.create_sheet("Datos")
d.column_dimensions["A"].width = 46; d.column_dimensions["B"].width = 20; d.column_dimensions["C"].width = 50
d["A1"] = "DATOS DEL CLIENTE"; d["A1"].font = Font(bold=True, size=13)
info = [("Empresa", "Comercializadora Ejemplo SAS"), ("NIT", "901000000-1"),
        ("Sector", "Comercio"), ("Régimen", "Ordinario"),
        ("Periodo analizado", "Ene–Dic 2025"), ("Fecha del diagnóstico", dt.date(2026, 10, 16)),
        ("Elaborado por", "Contador Público TP ____")]
for i, (k, v) in enumerate(info, 2):
    d.cell(i, 1, k); c = d.cell(i, 2, v); c.font = INPUT
d["B7"].number_format = "dd/mm/yyyy"
dvr = DataValidation(type="list", formula1='"SIMPLE,Ordinario,Persona natural,No responsable de IVA"')
d.add_data_validation(dvr); dvr.add("B5")

head(d, 10, ["ESTADO DE RESULTADOS (periodo)", "Valor COP", "Nota"])
er = [("Ventas netas", 1_200_000_000, ""), ("Costo de ventas", 840_000_000, ""),
      ("Gastos de administración", 150_000_000, ""), ("Gastos de ventas", 95_000_000, ""),
      ("Gastos financieros (intereses)", 38_000_000, ""), ("Otros ingresos (gastos) netos", 0, "Positivo = ingreso"),
      ("Impuesto de renta", 25_000_000, "")]
for i, (k, v, n) in enumerate(er, 11):
    d.cell(i, 1, k); c = d.cell(i, 2, v); c.font = INPUT; c.number_format = COP; d.cell(i, 3, n)

head(d, 19, ["BALANCE (fecha de corte)", "Valor COP", "Nota"])
bal = [("Efectivo y equivalentes", 35_000_000, ""), ("Cuentas por cobrar clientes", 260_000_000, ""),
       ("Inventarios", 180_000_000, ""), ("Otros activos corrientes", 20_000_000, ""),
       ("Activos no corrientes (propiedad, planta, otros)", 210_000_000, ""),
       ("Proveedores", 95_000_000, ""), ("Obligaciones financieras corto plazo", 140_000_000, ""),
       ("Impuestos por pagar", 30_000_000, ""), ("Otros pasivos corrientes", 25_000_000, ""),
       ("Pasivos no corrientes", 120_000_000, "")]
for i, (k, v, n) in enumerate(bal, 20):
    d.cell(i, 1, k); c = d.cell(i, 2, v); c.font = INPUT; c.number_format = COP; d.cell(i, 3, n)
d["A31"] = "Gasto fijo mensual promedio (nómina, arriendo, servicios…)"
d["B31"] = 22_000_000; d["B31"].font = INPUT; d["B31"].number_format = COP
d["C31"] = "Para calcular los meses de caja disponibles"

head(d, 33, ["CHEQUEO TRIBUTARIO", "Sí / No / N/A", "Riesgo si es 'No'"])
chk = [("RUT actualizado (responsabilidades y actividad)", "Sí", "Sanciones / responsabilidades mal asignadas"),
       ("Facturación electrónica habilitada y emitiendo", "Sí", "Sanción por no facturar; rechazo de costos del cliente"),
       ("Nómina electrónica transmitida mensualmente", "No", "Costos de nómina no deducibles"),
       ("Documento soporte en compras a no obligados", "No", "Costos y gastos no deducibles"),
       ("Retenciones en la fuente declaradas y pagadas a tiempo", "Sí", "Ineficacia de la declaración + intereses"),
       ("IVA / impuesto al consumo declarado a tiempo", "Sí", "Sanción por extemporaneidad e intereses"),
       ("Renta del último año presentada", "Sí", "Omisión / emplazamiento DIAN"),
       ("Información exógena reportada", "No", "Sanción por no enviar información"),
       ("ICA / industria y comercio al día", "Sí", "Sanciones municipales"),
       ("Contabilidad al día (cierre mensual ≤ 30 días)", "No", "Decisiones con información atrasada")]
dv = DataValidation(type="list", formula1='"Sí,No,N/A"', allow_blank=True)
d.add_data_validation(dv)
for i, (k, v, r) in enumerate(chk, 34):
    d.cell(i, 1, k); c = d.cell(i, 2, v); c.font = INPUT; c.alignment = Alignment(horizontal="center")
    d.cell(i, 3, r); dv.add(c.coordinate)
for r in list(range(11, 18)) + list(range(20, 30)) + [31] + list(range(34, 44)):
    for col in "ABC": d[f"{col}{r}"].border = B

# ---------- Parametros ----------
p = wb.create_sheet("Parametros")
head(p, 1, ["Indicador", "Dirección", "Verde si", "Rojo si", "Acción si ROJO", "Acción si AMARILLO", "Tipo"])
for col, w in zip("ABCDEF", [28, 12, 10, 10, 60, 60]):
    p.column_dimensions[col].width = w
# (nombre, dirección: 1 = más alto es mejor, -1 = más bajo es mejor, verde, rojo, formato, msg rojo, msg amarillo)
KPIS = [
    ("Margen bruto", 1, 0.35, 0.20, "0.0%", "Revisar precios, descuentos y costo de compra; el margen no cubre la operación.", "Negociar con proveedores y revisar la mezcla de productos."),
    ("Margen operacional", 1, 0.10, 0.03, "0.0%", "Los gastos consumen el margen: recortar gastos fijos y revisar la estructura.", "Controlar los gastos de administración y ventas mes a mes."),
    ("Margen neto", 1, 0.07, 0.02, "0.0%", "La rentabilidad final es crítica: revisar la carga financiera y tributaria.", "Planeación tributaria y reducción del costo financiero."),
    ("Razón corriente", 1, 1.5, 1.0, "0.00", "Riesgo de iliquidez: los pasivos de corto plazo superan los activos corrientes.", "Mejorar el capital de trabajo: cobrar más rápido y reestructurar la deuda de corto plazo."),
    ("Prueba ácida", 1, 1.0, 0.7, "0.00", "Sin inventario no alcanza a cubrir las deudas de corto plazo.", "Reducir el inventario inmovilizado."),
    ("Días de cartera", -1, 45, 90, "0", "Los clientes pagan muy tarde: política de crédito, cobro y descuentos por pronto pago.", "Seguimiento semanal de la cartera por edades."),
    ("Días de inventario", -1, 60, 120, "0", "Exceso de inventario: liquidar las referencias de baja rotación.", "Ajustar las compras a la rotación real."),
    ("Días de proveedores", 1, 30, 15, "0", "Pagas demasiado rápido: negociar plazos con los proveedores.", "Buscar más plazo con los proveedores clave."),
    ("Ciclo de caja (días)", -1, 45, 90, "0", "El negocio financia a los clientes durante demasiado tiempo: actuar sobre cartera, inventario y proveedores.", "Monitorear el ciclo de caja mensualmente."),
    ("Endeudamiento", -1, 0.50, 0.70, "0.0%", "Deuda alta: frenar el endeudamiento nuevo y reestructurar.", "Vigilar el nivel de deuda antes de financiar crecimiento."),
    ("Cobertura de intereses", 1, 3.0, 1.5, "0.0", "La utilidad operativa apenas cubre los intereses: refinanciar.", "Buscar tasas más bajas o prepagar la deuda cara."),
    ("Meses de caja", 1, 3.0, 1.0, "0.0", "Menos de 1 mes de caja: plan de tesorería urgente (flujo a 13 semanas).", "Construir una reserva de al menos 3 meses de gastos fijos."),
    ("Cumplimiento tributario", 1, 1.0, 0.8, "0%", "Riesgo alto de sanciones DIAN: normalizar las obligaciones pendientes de inmediato.", "Corregir los puntos pendientes del chequeo tributario."),
]
for i, (n, dr, g, r, fmt, mr, ma) in enumerate(KPIS, 2):
    p.cell(i, 1, n); p.cell(i, 2, dr).font = INPUT
    for col, v in ((3, g), (4, r)):
        c = p.cell(i, col, v); c.font = INPUT; c.number_format = fmt
    p.cell(i, 5, mr).alignment = WRAP; p.cell(i, 6, ma).alignment = WRAP
    p.cell(i, 7, "pct" if "%" in fmt else "num")
    for col in range(1, 7): p.cell(i, col).border = B
p["A17"] = "Dirección: 1 = más alto es mejor · -1 = más bajo es mejor. Los rangos son referencias generales; ajústalos por sector."
p["A17"].font = Font(italic=True)

# ---------- Diagnostico ----------
g = wb.create_sheet("Diagnostico", 1)
for col, w in zip("ABCDEF", [26, 14, 11, 11, 12, 58]):
    g.column_dimensions[col].width = w
g.merge_cells("A1:F1"); g["A1"] = "DIAGNÓSTICO FINANCIERO"
g["A1"].font = Font(bold=True, size=16, color="1F3864")
g["A2"] = '="Empresa: "&Datos!B2&"   ·   NIT: "&Datos!B3&"   ·   Sector: "&Datos!B4'
g["A3"] = '="Periodo: "&Datos!B6&"   ·   Régimen: "&Datos!B5'
g.merge_cells("A2:F2"); g.merge_cells("A3:D3")
g["E3"] = "Fecha:"; g["E3"].alignment = Alignment(horizontal="right")
g["F3"] = "=Datos!B7"; g["F3"].number_format = "dd/mm/yyyy"; g["F3"].alignment = Alignment(horizontal="left")

# Resumen de cifras
g["A5"] = "CIFRAS CLAVE"; g["A5"].font = NEG
V = "Datos!B11"; CV = "Datos!B12"
cif = [("Ventas", f"={V}"),
       ("Utilidad bruta", f"={V}-{CV}"),
       ("Utilidad operacional (EBIT)", f"={V}-{CV}-Datos!B13-Datos!B14"),
       ("Utilidad neta", f"={V}-{CV}-Datos!B13-Datos!B14-Datos!B15+Datos!B16-Datos!B17"),
       ("Activo total", "=SUM(Datos!B20:B24)"),
       ("Pasivo total", "=SUM(Datos!B25:B29)"),
       ("Patrimonio", "=B10-B11")]
for i, (k, f_) in enumerate(cif, 6):
    g.cell(i, 1, k); c = g.cell(i, 2, f_); c.number_format = COP
    for col in (1, 2): g.cell(i, col).border = B
AC = "SUM(Datos!B20:B23)"; PC = "SUM(Datos!B25:B28)"

# Indicadores
head(g, 14, ["Indicador", "Resultado", "Verde si", "Rojo si", "Semáforo", "Acción recomendada"])
calc = [
    f"=IFERROR(B7/{V},\"N/D\")",
    f"=IFERROR(B8/{V},\"N/D\")",
    f"=IFERROR(B9/{V},\"N/D\")",
    f"=IFERROR({AC}/{PC},\"N/D\")",
    f"=IFERROR(({AC}-Datos!B22)/{PC},\"N/D\")",
    f"=IFERROR(Datos!B21/{V}*365,\"N/D\")",
    f"=IFERROR(Datos!B22/{CV}*365,\"N/D\")",
    f"=IFERROR(Datos!B25/{CV}*365,\"N/D\")",
    "=IFERROR(B20+B21-B22,\"N/D\")",
    "=IFERROR(B11/B10,\"N/D\")",
    "=IFERROR(B8/Datos!B15,\"N/D\")",
    "=IFERROR(Datos!B20/Datos!B31,\"N/D\")",
    '=IFERROR(COUNTIF(Datos!B34:B43,"Sí")/(COUNTIF(Datos!B34:B43,"Sí")+COUNTIF(Datos!B34:B43,"No")),"N/D")',
]
for k, ((n, *_rest), fx) in enumerate(zip(KPIS, calc)):
    r, pr = 15 + k, 2 + k
    fmt = KPIS[k][4]
    g.cell(r, 1, f"=Parametros!A{pr}")
    c = g.cell(r, 2, fx); c.number_format = fmt
    for col, src in ((3, "C"), (4, "D")):
        c = g.cell(r, col, f"=Parametros!{src}{pr}"); c.number_format = fmt
    g.cell(r, 5, (f'=IF(ISNUMBER(B{r}),IF(Parametros!B{pr}=1,'
                  f'IF(B{r}>=C{r},"Verde",IF(B{r}<D{r},"Rojo","Amarillo")),'
                  f'IF(B{r}<=C{r},"Verde",IF(B{r}>D{r},"Rojo","Amarillo"))),"N/D")'))
    g.cell(r, 6, f'=IF(E{r}="Rojo",Parametros!E{pr},IF(E{r}="Amarillo",Parametros!F{pr},IF(E{r}="Verde","Mantener.","Falta información.")))')
    g.cell(r, 5).alignment = Alignment(horizontal="center")
    g.cell(r, 6).alignment = WRAP
    for col in range(1, 7): g.cell(r, col).border = B
last = 14 + len(KPIS)
rng = f"E15:E{last}"
g.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Verde"'], fill=PatternFill("solid", fgColor="C6EFCE"), font=Font(color="006100", bold=True)))
g.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Amarillo"'], fill=PatternFill("solid", fgColor="FFEB9C"), font=Font(color="9C5700", bold=True)))
g.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Rojo"'], fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(color="9C0006", bold=True)))

# Puntaje
s = last + 2
g.cell(s, 1, "PUNTAJE DE SALUD FINANCIERA").font = NEG
g.cell(s, 2, f'=IFERROR((COUNTIF({rng},"Verde")*2+COUNTIF({rng},"Amarillo"))/((COUNTIF({rng},"Verde")+COUNTIF({rng},"Amarillo")+COUNTIF({rng},"Rojo"))*2),0)').number_format = "0%"
g.cell(s, 2).font = Font(bold=True, size=14)
g.cell(s, 3, f'=IF(B{s}>=0.75,"SALUDABLE",IF(B{s}>=0.5,"EN ALERTA","CRÍTICO"))').font = Font(bold=True, size=12)
g.merge_cells(start_row=s, start_column=3, end_row=s, end_column=4)
g.cell(s + 1, 1, "Indicadores en rojo / amarillo / verde")
g.cell(s + 1, 2, f'=COUNTIF({rng},"Rojo")&" / "&COUNTIF({rng},"Amarillo")&" / "&COUNTIF({rng},"Verde")')
g.conditional_formatting.add(f"C{s}", FormulaRule(formula=[f'$C${s}="CRÍTICO"'], fill=PatternFill("solid", fgColor="FFC7CE")))
g.conditional_formatting.add(f"C{s}", FormulaRule(formula=[f'$C${s}="EN ALERTA"'], fill=PatternFill("solid", fgColor="FFEB9C")))
g.conditional_formatting.add(f"C{s}", FormulaRule(formula=[f'$C${s}="SALUDABLE"'], fill=PatternFill("solid", fgColor="C6EFCE")))

# Próximo paso
n = s + 3
g.cell(n, 1, "PRÓXIMO PASO RECOMENDADO").font = NEG
g.merge_cells(start_row=n + 1, start_column=1, end_row=n + 2, end_column=6)
g.cell(n + 1, 1, (f'=IF(COUNTIF({rng},"Rojo")>=3,"Acompañamiento financiero mensual: hay "&COUNTIF({rng},"Rojo")&" indicadores críticos que requieren seguimiento continuo.",'
                  f'IF(COUNTIF({rng},"Rojo")>=1,"Plan de acción de 90 días sobre los indicadores en rojo, con revisión mensual.",'
                  f'"Seguimiento trimestral y planeación tributaria para proteger la rentabilidad."))'))
g.cell(n + 1, 1).alignment = WRAP; g.cell(n + 1, 1).fill = GRIS
g.cell(n + 4, 1, "Diagnóstico basado en la información suministrada por el cliente, sin auditar. Rangos de referencia generales.").font = Font(italic=True, size=8)

g.print_area = f"A1:F{n + 4}"
g.page_setup.orientation = "portrait"; g.page_setup.fitToWidth = 1; g.page_setup.fitToHeight = 1
g.sheet_properties.pageSetUpPr.fitToPage = True
g.print_options.horizontalCentered = True
g.page_margins.left = g.page_margins.right = 0.4

# ---------- Orden de hallazgos (columna auxiliar fuera del área de impresión) ----------
g["G14"] = "Orden"; g["G14"].font = Font(color="808080", italic=True)
for r in range(15, last + 1):
    c = g.cell(r, 7, (f'=IF(E{r}="Rojo",COUNTIF(E$15:E{r},"Rojo"),IF(E{r}="Amarillo",'
                      f'COUNTIF($E$15:$E${last},"Rojo")+COUNTIF(E$15:E{r},"Amarillo"),""))'))
    c.font = Font(color="808080")

# ---------- Parámetros de la propuesta ----------
head(p, 20, ["Plan", "Precio mensual", "Incluye", "", "", ""])
PLANES = [("Básico", 450_000, "Contabilidad mensual, declaraciones tributarias y facturación electrónica."),
          ("Empresa", 750_000, "Básico + nómina electrónica, retenciones, información exógena e informe mensual."),
          ("Crecimiento", 1_100_000, "Empresa + flujo de caja, presupuesto, indicadores y reunión mensual de resultados."),
          ("CFO fraccional", 2_200_000, "Crecimiento + flujo de caja a 13 semanas, gestión de cartera e inventario, análisis de precios y márgenes, y acompañamiento en decisiones financieras.")]
for i, (n_, pr_, inc) in enumerate(PLANES, 21):
    p.cell(i, 1, n_).border = B
    c = p.cell(i, 2, pr_); c.font = INPUT; c.number_format = COP; c.border = B
    c = p.cell(i, 3, inc); c.font = INPUT; c.alignment = WRAP
    p.merge_cells(start_row=i, start_column=3, end_row=i, end_column=6); p.row_dimensions[i].height = 30
head(p, 27, ["Datos del contador", "", "", "", "", ""])
CONT = [("Nombre", "[Tu nombre]"), ("Tarjeta profesional", "TP ______"), ("Teléfono / WhatsApp", "[300 000 0000]"),
        ("Correo", "[correo@dominio.com]"), ("Ciudad", "Montería"), ("Validez de la propuesta (días)", 15),
        ("Condiciones de pago", "Pago anticipado del día 1 al 5 de cada mes. Atraso mayor de 15 días suspende el servicio. Valores antes de IVA, si aplica.")]
for i, (k, v) in enumerate(CONT, 28):
    p.cell(i, 1, k); c = p.cell(i, 2, v); c.font = INPUT
p.merge_cells("B34:F34"); p["B34"].alignment = WRAP; p.row_dimensions[34].height = 30
d["A9"] = "Plan elegido (opcional; vacío = recomendado)"
dvp = DataValidation(type="list", formula1="=Parametros!$A$21:$A$24", allow_blank=True)
d.add_data_validation(dvp); dvp.add("B9"); d["B9"].font = INPUT

# ---------- Propuesta ----------
q = wb.create_sheet("Propuesta", 2)
for col, w in zip("ABCDEF", [26, 16, 14, 14, 14, 48]):
    q.column_dimensions[col].width = w
q.merge_cells("A1:F1"); q["A1"] = "PROPUESTA DE SERVICIOS CONTABLES Y FINANCIEROS"
q["A1"].font = Font(bold=True, size=15, color="1F3864")
q.merge_cells("A2:F2"); q["A2"] = '="Para: "&Datos!B2&"   ·   NIT: "&Datos!B3'
q["A3"] = "Fecha:"; q["B3"] = "=Datos!B7"; q["C3"] = "Válida hasta:"; q["D3"] = "=Datos!B7+Parametros!B33"
q["E3"] = "=Parametros!B32"
for c in ("B3", "D3"): q[c].number_format = "dd/mm/yyyy"; q[c].alignment = Alignment(horizontal="left")
q["C3"].alignment = Alignment(horizontal="right")


def titulo(r, t):
    q.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    c = q.cell(r, 1, t); c.fill, c.font = AZUL, BL


titulo(5, "1. SITUACIÓN ENCONTRADA")
q["A6"] = "Salud financiera"; q["B6"] = f"=Diagnostico!B{s}"; q["B6"].number_format = "0%"; q["B6"].font = NEG
q["C6"] = f"=Diagnostico!C{s}"; q["C6"].font = NEG
q["D6"] = "Rojo / amarillo / verde"; q["F6"] = f"=Diagnostico!B{s + 1}"
head(q, 7, ["Hallazgo prioritario", "Resultado", "Meta", "Estado", "", "Qué haremos"])
q.merge_cells("D7:E7")
for i in range(1, 6):
    r = 7 + i
    m = f'MATCH({i},Diagnostico!$G$15:$G${last},0)'
    q.cell(r, 1, f'=IFERROR(INDEX(Diagnostico!$A$15:$A${last},{m}),"")')
    q.cell(r, 2, f'=IFERROR(INDEX(Diagnostico!$B$15:$B${last},{m}),"")')
    q.cell(r, 3, f'=IFERROR(INDEX(Diagnostico!$C$15:$C${last},{m}),"")')
    q.cell(r, 4, f'=IFERROR(INDEX(Diagnostico!$E$15:$E${last},{m}),"")')
    q.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
    q.cell(r, 6, f'=IFERROR(INDEX(Diagnostico!$F$15:$F${last},{m}),"")').alignment = WRAP
    q.cell(r, 2).number_format = "0.0"; q.cell(r, 3).number_format = "0.0"
    c = q.cell(r, 7, f'=IFERROR(INDEX(Parametros!$G$2:$G${len(KPIS) + 1},{m}),"")'); c.font = Font(color="808080")
    q.cell(r, 4).alignment = Alignment(horizontal="center")
    q.row_dimensions[r].height = 32
    for col in range(1, 7): q.cell(r, col).border = B
q.conditional_formatting.add("B8:C12", Rule(type="expression", formula=['$G8="pct"'], dxf=DifferentialStyle(numFmt=NumberFormat(numFmtId=170, formatCode="0.0%"))))
q.conditional_formatting.add("D8:D12", CellIsRule(operator="equal", formula=['"Rojo"'], fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(color="9C0006", bold=True)))
q.conditional_formatting.add("D8:D12", CellIsRule(operator="equal", formula=['"Amarillo"'], fill=PatternFill("solid", fgColor="FFEB9C"), font=Font(color="9C5700", bold=True)))

titulo(14, "2. CAJA QUE SE PUEDE LIBERAR (estimación)")
q["A15"] = "Cartera por encima de la meta"
q["B15"] = '=IF(ISNUMBER(Diagnostico!B20),MAX(0,Diagnostico!B20-Diagnostico!C20)*Datos!B11/365,0)'
q["A16"] = "Inventario por encima de la meta"
q["B16"] = '=IF(ISNUMBER(Diagnostico!B21),MAX(0,Diagnostico!B21-Diagnostico!C21)*Datos!B12/365,0)'
q["A17"] = "Total caja potencial"; q["A17"].font = NEG
q["B17"] = "=B15+B16"; q["B17"].font = Font(bold=True, size=12)
for r in (15, 16, 17): q[f"B{r}"].number_format = COP
q.merge_cells("C15:F17")
q["C15"] = ("Dinero inmovilizado en cartera e inventario por encima de los días meta. Llevar esos días a la meta "
            "libera esa caja sin endeudarse. Es una estimación basada en las cifras del periodo; el resultado real "
            "depende de la ejecución.")
q["C15"].alignment = WRAP; q["C15"].font = Font(italic=True, size=9)

titulo(19, "3. PLAN RECOMENDADO")
q["A20"] = "Plan sugerido por el diagnóstico"
q["B20"] = (f'=IF(COUNTIF(Diagnostico!$E$15:$E${last},"Rojo")>=3,"CFO fraccional",'
            f'IF(COUNTIF(Diagnostico!$E$15:$E${last},"Rojo")>=1,"Crecimiento","Empresa"))')
q["A21"] = "Plan propuesto"; q["A21"].font = NEG
q["B21"] = '=IF(Datos!B9<>"",Datos!B9,B20)'; q["B21"].font = Font(bold=True, size=12, color="1F3864")
q["A22"] = "Incluye"
q.merge_cells("B22:F23")
q["B22"] = '=IFERROR(VLOOKUP(B21,Parametros!$A$21:$C$24,3,FALSE),"")'; q["B22"].alignment = WRAP

titulo(25, "4. CÓMO TRABAJAREMOS")
pasos = [("Semana 1", "Recepción de información y accesos al software contable; plan de normalización de los hallazgos en rojo."),
         ("Semanas 2–4", "Puesta al día contable y tributaria; primeras acciones sobre cartera, inventario y caja."),
         ("Cada mes", "Cierre contable, declaraciones a tiempo e informe con indicadores y semáforo (este mismo formato)."),
         ("Mes 3", "Revisión de resultados frente a este diagnóstico: ¿cuánta caja se liberó y qué indicadores mejoraron?")]
for i, (k, v) in enumerate(pasos, 26):
    q.cell(i, 1, k).font = NEG
    q.merge_cells(start_row=i, start_column=2, end_row=i, end_column=6)
    q.cell(i, 2, v).alignment = WRAP; q.row_dimensions[i].height = 28

titulo(31, "5. INVERSIÓN")
q["A32"] = "Valor mensual"; q["A32"].font = NEG
q["B32"] = '=IFERROR(VLOOKUP(B21,Parametros!$A$21:$B$24,2,FALSE),"")'
q["B32"].number_format = COP; q["B32"].font = Font(bold=True, size=12)
q["C32"] = '=IF(AND(ISNUMBER(B32),B17>0),B32*12/B17,"")'; q["C32"].number_format = "0.0%"; q["C32"].font = Font(bold=True)
q.merge_cells("D32:F32"); q["D32"] = "del valor anual frente a la caja potencial a liberar"; q["D32"].font = Font(italic=True, size=9)
q["A33"] = "Condiciones"
q.merge_cells("B33:F33"); q["B33"] = "=Parametros!B34"; q["B33"].alignment = WRAP; q.row_dimensions[33].height = 30

titulo(35, "6. PRÓXIMO PASO")
q.merge_cells("A36:F37")
q["A36"] = ("Para iniciar: (1) aceptar esta propuesta respondiendo este mensaje o firmando abajo, (2) realizar el primer pago "
            "y (3) enviar la información contable y los accesos. Iniciamos en los 3 días hábiles siguientes.")
q["A36"].alignment = WRAP

q["A40"] = "_____________________________"; q["D40"] = "_____________________________"
q["A41"] = "=Parametros!B28"; q["D41"] = "Aceptado por: " ; q["A42"] = '="Contador Público · "&Parametros!B29'
q["D42"] = "=Datos!B2"
q["A43"] = '=Parametros!B30&" · "&Parametros!B31'; q["A43"].font = Font(size=9)

q.print_area = "A1:F43"
q.page_setup.orientation = "portrait"; q.page_setup.fitToWidth = 1; q.page_setup.fitToHeight = 1
q.sheet_properties.pageSetUpPr.fitToPage = True
q.print_options.horizontalCentered = True
q.page_margins.left = q.page_margins.right = 0.4

wb.save("plantillas/diagnostico-financiero-pyme.xlsx")
