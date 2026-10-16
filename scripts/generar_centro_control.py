"""Genera 'centro-de-control.xlsx': CRM, clientes, obligaciones, horas, facturación y tablero."""
import datetime as dt
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

AZUL = PatternFill("solid", fgColor="1F3864")
INPUT = Font(color="0000FF")
BL = Font(color="FFFFFF", bold=True)
NEG = Font(bold=True)
B = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
COP = '#,##0;[Red]-#,##0'
FECHA = "dd/mm/yyyy"
N = 300  # filas por tabla
D = dt.date

ROJO = (PatternFill("solid", fgColor="FFC7CE"), Font(color="9C0006", bold=True))
AMAR = (PatternFill("solid", fgColor="FFEB9C"), Font(color="9C5700", bold=True))
VERDE = (PatternFill("solid", fgColor="C6EFCE"), Font(color="006100", bold=True))


def cf(ws, rng, val, style):
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{val}"'], fill=style[0], font=style[1]))


def tabla(ws, headers, widths):
    for j, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(1, j, h); c.fill, c.font = AZUL, BL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(headers)).column_letter}{N + 1}"


def lista(ws, rng, src):
    dv = DataValidation(type="list", formula1=src, allow_blank=True)
    ws.add_data_validation(dv); dv.add(rng)


wb = Workbook()

# ---------------- Instrucciones ----------------
ins = wb.active; ins.title = "Instrucciones"
for i, (t, b) in enumerate([
    ("CENTRO DE CONTROL DE LA FIRMA", True), ("", False),
    ("Tablero: indicadores del negocio con semáforo frente a las metas. Cambia el 'Mes de análisis' en Parametros.", False),
    ("Pipeline: cada prospecto y su etapa. Regla: toda oportunidad abierta tiene un próximo paso con fecha.", False),
    ("Clientes: clientes activos, plan, valor mensual, riesgo y rentabilidad (horas del mes).", False),
    ("Obligaciones: TODOS los vencimientos de TODOS los clientes. Cárgalos cada año desde el calendario oficial DIAN/municipal.", False),
    ("Horas: registro diario de tiempo por cliente (base del ingreso por hora).", False),
    ("Facturacion: facturas emitidas, pagos y cartera de la firma por edades.", False),
    ("Parametros: planes, precios, listas desplegables y metas.", False), ("", False),
    ("Celdas en azul = datos que llenas. El resto se calcula solo. Los datos de ejemplo son ficticios: bórralos al empezar.", True),
], 1):
    ins.cell(i, 1, t).font = Font(bold=b, size=14 if i == 1 else 11)
ins.column_dimensions["A"].width = 120

# ---------------- Parametros ----------------
p = wb.create_sheet("Parametros")
p["A1"] = "Mes de análisis (día 1)"; p["A1"].font = NEG
p["B1"] = D(2026, 10, 1); p["B1"].number_format = FECHA; p["B1"].font = INPUT
p["A2"] = "Fecha de hoy"; p["B2"] = "=TODAY()"; p["B2"].number_format = FECHA
LISTAS = {
    "Plan": ["Básico", "Empresa", "Crecimiento", "CFO fraccional", "Renta PN", "Proyecto"],
    "Segmento": ["Pyme comercio", "Pyme servicios", "Emprendedor digital", "Persona natural", "Agro / ganadero", "Otro"],
    "Canal": ["Referido", "Red personal", "Google", "WhatsApp", "LinkedIn", "Alianza", "Charla / taller", "Pauta"],
    "Etapa": ["1 Prospecto", "2 Contactado", "3 Diagnóstico agendado", "4 Diagnóstico entregado",
              "5 Propuesta enviada", "6 Negociación", "7 Ganado", "8 Perdido"],
    "Obligación": ["IVA", "Retención en la fuente", "Renta PJ", "Renta PN", "SIMPLE anticipo", "SIMPLE anual",
                   "ICA", "Exógena", "Nómina electrónica", "PILA", "Estados financieros", "Renovación matrícula mercantil", "Otra"],
    "Estado obligación": ["Pendiente", "En preparación", "Enviada al cliente", "Aprobada", "Presentada"],
    "Actividad": ["Contabilidad", "Tributario", "Nómina", "Informe / reunión", "Diagnóstico", "Comercial", "Administrativo"],
    "Estado cliente": ["Activo", "Suspendido", "Retirado"],
    "Riesgo": ["Bajo", "Medio", "Alto"],
    "Motivo pérdida": ["Precio", "Ya tiene contador", "Sin respuesta", "No es el momento", "Competencia", "No califica"],
}
col = 4
REF = {}
for nombre, vals in LISTAS.items():
    c = p.cell(1, col, nombre); c.fill, c.font = AZUL, BL
    for i, v in enumerate(vals, 2):
        p.cell(i, col, v).font = INPUT
    L = c.column_letter
    REF[nombre] = f"=Parametros!${L}$2:${L}${len(vals) + 1}"
    p.column_dimensions[L].width = 24
    col += 1
p.column_dimensions["A"].width = 34; p.column_dimensions["B"].width = 16
# precios
p["A4"] = "Plan"; p["B4"] = "Precio mensual"
for c in ("A4", "B4"): p[c].fill, p[c].font = AZUL, BL
for i, (pl, pr) in enumerate([("Básico", 450_000), ("Empresa", 750_000), ("Crecimiento", 1_100_000),
                              ("CFO fraccional", 2_200_000), ("Renta PN", 250_000), ("Proyecto", 0)], 5):
    p.cell(i, 1, pl); c = p.cell(i, 2, pr); c.font = INPUT; c.number_format = COP
PRECIOS = "Parametros!$A$5:$B$10"
# metas
p["A13"] = "Meta"; p["B13"] = "Valor"
for c in ("A13", "B13"): p[c].fill, p[c].font = AZUL, BL
METAS = [("MRR meta (COP/mes)", 15_000_000, COP), ("Clientes nuevos / mes", 3, "0"),
         ("Tasa de cierre mínima", 0.30, "0%"), ("Cartera > 30 días máx. (% del MRR)", 0.05, "0%"),
         ("Ingreso por hora mínimo", 80_000, COP), ("Concentración máx. por cliente", 0.20, "0%"),
         ("Horas facturables disponibles / mes", 140, "0")]
for i, (k, v, f) in enumerate(METAS, 14):
    p.cell(i, 1, k); c = p.cell(i, 2, v); c.font = INPUT; c.number_format = f

# ---------------- Pipeline ----------------
pl = wb.create_sheet("Pipeline")
tabla(pl, ["Fecha", "Prospecto", "Segmento", "Canal", "Contacto (tel/correo)", "Etapa", "Plan propuesto",
           "Valor mensual", "Próximo paso", "Fecha próximo paso", "Alerta", "Motivo pérdida", "Abierta (1/0)"],
      [11, 30, 18, 14, 22, 22, 15, 14, 30, 13, 12, 18, 9])
for r in range(2, N + 2):
    pl[f"H{r}"] = f'=IF(G{r}="","",IFERROR(VLOOKUP(G{r},{PRECIOS},2,FALSE),0))'
    pl[f"K{r}"] = (f'=IF(OR(B{r}="",F{r}="7 Ganado",F{r}="8 Perdido"),"",'
                   f'IF(J{r}="","SIN FECHA",IF(J{r}<Parametros!$B$2,"VENCIDO","OK")))')
    pl[f"M{r}"] = f'=IF(AND(B{r}<>"",F{r}<>"7 Ganado",F{r}<>"8 Perdido"),1,0)'
    pl[f"A{r}"].number_format = FECHA; pl[f"J{r}"].number_format = FECHA; pl[f"H{r}"].number_format = COP
    for c in "ABCDEFGIJL": pl[f"{c}{r}"].font = INPUT
lista(pl, f"C2:C{N + 1}", REF["Segmento"]); lista(pl, f"D2:D{N + 1}", REF["Canal"])
lista(pl, f"F2:F{N + 1}", REF["Etapa"]); lista(pl, f"G2:G{N + 1}", REF["Plan"])
lista(pl, f"L2:L{N + 1}", REF["Motivo pérdida"])
cf(pl, f"K2:K{N + 1}", "VENCIDO", ROJO); cf(pl, f"K2:K{N + 1}", "SIN FECHA", AMAR); cf(pl, f"K2:K{N + 1}", "OK", VERDE)
ej_pl = [(D(2026, 10, 1), "Ferretería El Progreso", "Pyme comercio", "Referido", "300 000 0001", "5 Propuesta enviada", "Crecimiento", "Llamar para resolver dudas", D(2026, 10, 20), ""),
         (D(2026, 10, 2), "Tienda Online Moda Sinú", "Emprendedor digital", "WhatsApp", "300 000 0002", "3 Diagnóstico agendado", "Básico", "Reunión de diagnóstico", D(2026, 10, 17), ""),
         (D(2026, 10, 3), "Transportes del Caribe SAS", "Pyme servicios", "Alianza", "300 000 0003", "2 Contactado", "Empresa", "Enviar mensaje de seguimiento", D(2026, 10, 10), ""),
         (D(2026, 9, 20), "Restaurante La Ribera", "Pyme servicios", "Red personal", "300 000 0004", "7 Ganado", "Empresa", "", None, ""),
         (D(2026, 9, 22), "Distribuidora Andina", "Pyme comercio", "Google", "300 000 0005", "8 Perdido", "Empresa", "", None, "Precio"),
         (D(2026, 10, 5), "Hacienda Los Robles", "Agro / ganadero", "Red personal", "300 000 0006", "4 Diagnóstico entregado", "CFO fraccional", "Presentar propuesta", D(2026, 10, 21), "")]
for i, row in enumerate(ej_pl, 2):
    for j, v in zip("ABCDEFGIJL", row):
        if v not in (None, ""): pl[f"{j}{i}"] = v

# ---------------- Clientes ----------------
cl = wb.create_sheet("Clientes")
tabla(cl, ["Cliente", "NIT / C.C.", "Segmento", "Plan", "Valor mensual", "Fecha inicio", "Estado", "Riesgo",
           "Satisfacción (NPS 0–10)", "Horas mes análisis", "Ingreso por hora", "% del MRR", "Alerta rentabilidad"],
      [30, 14, 18, 15, 14, 12, 11, 9, 11, 11, 13, 9, 16])
for r in range(2, N + 2):
    cl[f"J{r}"] = f'=IF(A{r}="","",SUMIFS(Horas!$D$2:$D${N + 1},Horas!$B$2:$B${N + 1},A{r},Horas!$E$2:$E${N + 1},Parametros!$B$1))'
    cl[f"K{r}"] = f'=IF(OR(A{r}="",G{r}<>"Activo"),"",IF(N(J{r})=0,"Sin horas",E{r}/J{r}))'
    cl[f"L{r}"] = f'=IF(OR(A{r}="",G{r}<>"Activo"),"",IFERROR(E{r}/Tablero!$B$5,0))'
    cl[f"M{r}"] = (f'=IF(NOT(ISNUMBER(K{r})),"",IF(K{r}<Parametros!$B$18,"BAJA",'
                   f'IF(K{r}<Parametros!$B$18*1.25,"AJUSTADA","SANA")))')
    cl[f"E{r}"].number_format = COP; cl[f"K{r}"].number_format = COP; cl[f"F{r}"].number_format = FECHA
    cl[f"L{r}"].number_format = "0%"; cl[f"J{r}"].number_format = "0.0"
    for c in "ABCDEFGHI": cl[f"{c}{r}"].font = INPUT
lista(cl, f"C2:C{N + 1}", REF["Segmento"]); lista(cl, f"D2:D{N + 1}", REF["Plan"])
lista(cl, f"G2:G{N + 1}", REF["Estado cliente"]); lista(cl, f"H2:H{N + 1}", REF["Riesgo"])
cf(cl, f"M2:M{N + 1}", "BAJA", ROJO); cf(cl, f"M2:M{N + 1}", "AJUSTADA", AMAR); cf(cl, f"M2:M{N + 1}", "SANA", VERDE)
cf(cl, f"H2:H{N + 1}", "Alto", ROJO); cf(cl, f"H2:H{N + 1}", "Medio", AMAR)
ej_cl = [("Restaurante La Ribera", "900100200", "Pyme servicios", "Empresa", 750_000, D(2026, 10, 1), "Activo", "Bajo", 9),
         ("Comercializadora Sinú SAS", "900300400", "Pyme comercio", "Crecimiento", 1_100_000, D(2026, 7, 1), "Activo", "Bajo", 10),
         ("Juan Pérez (independiente)", "1067000000", "Persona natural", "Básico", 450_000, D(2026, 8, 1), "Activo", "Medio", 7),
         ("Taller Mecánico Rápido", "901500600", "Pyme servicios", "Empresa", 750_000, D(2026, 5, 1), "Activo", "Alto", 5),
         ("Panadería Doña Ana", "901700800", "Pyme comercio", "Básico", 450_000, D(2026, 3, 1), "Retirado", "Bajo", 6)]
for i, row in enumerate(ej_cl, 2):
    for j, v in zip("ABCDEFGHI", row): cl[f"{j}{i}"] = v

# ---------------- Obligaciones ----------------
ob = wb.create_sheet("Obligaciones")
tabla(ob, ["Cliente", "Obligación", "Periodo", "Fecha de vencimiento", "Estado", "Responsable",
           "Días restantes", "Alerta", "Número formulario / soporte"],
      [30, 22, 14, 14, 18, 14, 10, 12, 24])
for r in range(2, N + 2):
    ob[f"G{r}"] = f'=IF(OR(A{r}="",D{r}=""),"",IF(E{r}="Presentada","",D{r}-Parametros!$B$2))'
    ob[f"H{r}"] = (f'=IF(OR(A{r}="",D{r}=""),"",IF(E{r}="Presentada","CUMPLIDA",'
                   f'IF(G{r}<0,"VENCIDA",IF(G{r}<=5,"URGENTE",IF(G{r}<=15,"PRÓXIMA","OK")))))')
    ob[f"D{r}"].number_format = FECHA
    for c in "ABCDEFI": ob[f"{c}{r}"].font = INPUT
lista(ob, f"B2:B{N + 1}", REF["Obligación"]); lista(ob, f"E2:E{N + 1}", REF["Estado obligación"])
cf(ob, f"H2:H{N + 1}", "VENCIDA", ROJO); cf(ob, f"H2:H{N + 1}", "URGENTE", ROJO)
cf(ob, f"H2:H{N + 1}", "PRÓXIMA", AMAR); cf(ob, f"H2:H{N + 1}", "CUMPLIDA", VERDE)
hoy = dt.date.today()
ej_ob = [("Restaurante La Ribera", "Retención en la fuente", "Sep 2026", hoy + dt.timedelta(days=3), "En preparación", "Titular"),
         ("Comercializadora Sinú SAS", "IVA", "Bim 4 2026", hoy + dt.timedelta(days=10), "Pendiente", "Titular"),
         ("Comercializadora Sinú SAS", "Retención en la fuente", "Sep 2026", hoy - dt.timedelta(days=2), "Presentada", "Titular"),
         ("Juan Pérez (independiente)", "Renta PN", "AG 2025", hoy + dt.timedelta(days=40), "Pendiente", "Titular"),
         ("Taller Mecánico Rápido", "Nómina electrónica", "Sep 2026", hoy - dt.timedelta(days=1), "Pendiente", "Titular")]
for i, row in enumerate(ej_ob, 2):
    for j, v in zip("ABCDEF", row): ob[f"{j}{i}"] = v
ob["J1"] = "Fechas de ejemplo calculadas a partir de la fecha de generación. Usa SIEMPRE el calendario oficial vigente (DIAN / municipio)."
ob["J1"].font = Font(italic=True, color="808080")

# ---------------- Horas ----------------
hr = wb.create_sheet("Horas")
tabla(hr, ["Fecha", "Cliente", "Actividad", "Horas", "Mes", "Nota"], [11, 30, 18, 8, 10, 40])
for r in range(2, N + 2):
    hr[f"E{r}"] = f'=IF(A{r}="","",DATE(YEAR(A{r}),MONTH(A{r}),1))'
    hr[f"A{r}"].number_format = FECHA; hr[f"E{r}"].number_format = "mmm-yy"; hr[f"D{r}"].number_format = "0.0"
    for c in "ABCDF": hr[f"{c}{r}"].font = INPUT
lista(hr, f"C2:C{N + 1}", REF["Actividad"])
lista(hr, f"B2:B{N + 1}", f"=Clientes!$A$2:$A${N + 1}")
ej_hr = [(D(2026, 10, 2), "Restaurante La Ribera", "Contabilidad", 4), (D(2026, 10, 6), "Restaurante La Ribera", "Nómina", 2),
         (D(2026, 10, 3), "Comercializadora Sinú SAS", "Contabilidad", 6), (D(2026, 10, 9), "Comercializadora Sinú SAS", "Informe / reunión", 3),
         (D(2026, 10, 7), "Juan Pérez (independiente)", "Tributario", 2),
         (D(2026, 10, 8), "Taller Mecánico Rápido", "Contabilidad", 7), (D(2026, 10, 12), "Taller Mecánico Rápido", "Tributario", 5),
         (D(2026, 9, 28), "Restaurante La Ribera", "Diagnóstico", 2)]
for i, row in enumerate(ej_hr, 2):
    for j, v in zip("ABCD", row): hr[f"{j}{i}"] = v

# ---------------- Facturacion ----------------
fa = wb.create_sheet("Facturacion")
tabla(fa, ["N.º factura", "Cliente", "Fecha emisión", "Fecha límite pago", "Valor", "Pagado", "Fecha pago",
           "Saldo", "Días de mora", "Estado"], [12, 30, 12, 12, 14, 14, 12, 14, 9, 12])
for r in range(2, N + 2):
    fa[f"D{r}"] = f'=IF(C{r}="","",C{r}+4)'
    fa[f"H{r}"] = f'=IF(E{r}="","",E{r}-N(F{r}))'
    fa[f"I{r}"] = f'=IF(OR(H{r}="",N(H{r})<=0),0,MAX(0,Parametros!$B$2-D{r}))'
    fa[f"J{r}"] = (f'=IF(E{r}="","",IF(H{r}<=0,"Pagada",IF(I{r}=0,"Al día",IF(I{r}<=15,"1-15",'
                   f'IF(I{r}<=30,"16-30","Más de 30")))))')
    for c in "CDG": fa[f"{c}{r}"].number_format = FECHA
    for c in "EFH": fa[f"{c}{r}"].number_format = COP
    for c in "ABCEFG": fa[f"{c}{r}"].font = INPUT
lista(fa, f"B2:B{N + 1}", f"=Clientes!$A$2:$A${N + 1}")
cf(fa, f"J2:J{N + 1}", "Más de 30", ROJO); cf(fa, f"J2:J{N + 1}", "16-30", AMAR); cf(fa, f"J2:J{N + 1}", "Pagada", VERDE)
ej_fa = [("FE-001", "Restaurante La Ribera", D(2026, 10, 1), 750_000, 750_000, D(2026, 10, 3)),
         ("FE-002", "Comercializadora Sinú SAS", D(2026, 10, 1), 1_100_000, 1_100_000, D(2026, 10, 2)),
         ("FE-003", "Juan Pérez (independiente)", D(2026, 10, 1), 450_000, None, None),
         ("FE-004", "Taller Mecánico Rápido", D(2026, 10, 1), 750_000, None, None),
         ("FE-005", "Taller Mecánico Rápido", hoy - dt.timedelta(days=45), 750_000, 300_000, hoy - dt.timedelta(days=20))]
for i, (n_, c_, f_, v_, pg, fp) in enumerate(ej_fa, 2):
    fa[f"A{i}"], fa[f"B{i}"], fa[f"C{i}"], fa[f"E{i}"] = n_, c_, f_, v_
    if pg: fa[f"F{i}"] = pg
    if fp: fa[f"G{i}"] = fp

# ---------------- Tablero ----------------
t = wb.create_sheet("Tablero", 1)
t.column_dimensions["A"].width = 40; t.column_dimensions["B"].width = 18
t.column_dimensions["C"].width = 16; t.column_dimensions["D"].width = 14
t["A1"] = "TABLERO DE LA FIRMA"; t["A1"].font = Font(bold=True, size=16, color="1F3864")
t["A2"] = "Mes de análisis:"; t["B2"] = "=Parametros!B1"; t["B2"].number_format = "mmm-yy"
for j, h in enumerate(["Indicador", "Valor", "Meta", "Estado"], 1):
    c = t.cell(4, j, h); c.fill, c.font = AZUL, BL
CLI = lambda c: f"Clientes!${c}$2:${c}${N + 1}"
PIP = lambda c: f"Pipeline!${c}$2:${c}${N + 1}"
OBL = lambda c: f"Obligaciones!${c}$2:${c}${N + 1}"
FAC = lambda c: f"Facturacion!${c}$2:${c}${N + 1}"
HOR = lambda c: f"Horas!${c}$2:${c}${N + 1}"
M1, M2 = "Parametros!$B$1", "EDATE(Parametros!$B$1,1)"
# (clave, nombre, fórmula, meta, formato, regla). {CLAVE} se reemplaza por la celda del indicador.
KPI = [
    (None, "FINANZAS", None, None, None, None),
    ("MRR", "Ingreso mensual recurrente (MRR)", f'=SUMIFS({CLI("E")},{CLI("G")},"Activo")', "=Parametros!B14", COP, "mayor"),
    ("ACT", "Clientes activos", f'=COUNTIFS({CLI("G")},"Activo")', None, "0", None),
    ("TIC", "Ticket promedio", '=IFERROR({MRR}/{ACT},0)', None, COP, None),
    ("FACM", "Facturado en el mes", f'=SUMIFS({FAC("E")},{FAC("C")},">="&{M1},{FAC("C")},"<"&{M2})', None, COP, None),
    ("CAR", "Cartera total por cobrar", f'=SUM({FAC("H")})', None, COP, None),
    ("C30", "Cartera con más de 30 días de mora", f'=SUMIFS({FAC("H")},{FAC("I")},">30")', None, COP, None),
    ("C30P", "Cartera > 30 días / MRR", '=IFERROR({C30}/{MRR},0)', "=Parametros!B17", "0.0%", "menor"),
    ("HRS", "Horas registradas en el mes", f'=SUMIFS({HOR("D")},{HOR("E")},{M1})', None, "0.0", None),
    ("IXH", "Ingreso por hora (MRR / horas)", '=IFERROR({MRR}/{HRS},0)', "=Parametros!B18", COP, "mayor"),
    ("CAP", "Uso de capacidad", '=IFERROR({HRS}/Parametros!B20,0)', 0.8, "0%", "menor"),
    ("CON", "Concentración: cliente más grande / MRR", f'=IFERROR(MAX({CLI("L")}),0)', "=Parametros!B19", "0%", "menor"),
    (None, "COMERCIAL", None, None, None, None),
    ("NUE", "Clientes nuevos en el mes", f'=COUNTIFS({CLI("F")},">="&{M1},{CLI("F")},"<"&{M2})', "=Parametros!B15", "0", "mayor"),
    ("OPA", "Oportunidades abiertas", f'=SUM({PIP("M")})', None, "0", None),
    ("PIPV", "Valor mensual del pipeline abierto", f'=SUMIFS({PIP("H")},{PIP("M")},1)', None, COP, None),
    ("CIE", "Tasa de cierre (ganados / cerrados)", f'=IFERROR(COUNTIF({PIP("F")},"7 Ganado")/(COUNTIF({PIP("F")},"7 Ganado")+COUNTIF({PIP("F")},"8 Perdido")),0)', "=Parametros!B16", "0%", "mayor"),
    ("PVE", "Oportunidades con próximo paso vencido o sin fecha", f'=COUNTIF({PIP("K")},"VENCIDO")+COUNTIF({PIP("K")},"SIN FECHA")', 0, "0", "menor"),
    (None, "OPERACIONES Y RIESGO", None, None, None, None),
    ("OBV", "Obligaciones VENCIDAS sin presentar", f'=COUNTIF({OBL("H")},"VENCIDA")', 0, "0", "menor"),
    ("OBU", "Obligaciones URGENTES (≤ 5 días)", f'=COUNTIF({OBL("H")},"URGENTE")', None, "0", None),
    ("OBP", "Obligaciones próximas (6–15 días)", f'=COUNTIF({OBL("H")},"PRÓXIMA")', None, "0", None),
    ("RAL", "Clientes activos con riesgo ALTO", f'=COUNTIFS({CLI("G")},"Activo",{CLI("H")},"Alto")', 0, "0", "menor"),
    ("RBA", "Clientes con rentabilidad BAJA", f'=COUNTIF({CLI("M")},"BAJA")', 0, "0", "menor"),
    (None, "CLIENTE", None, None, None, None),
    ("NPS", "NPS (promotores − detractores)", f'=IFERROR((COUNTIFS({CLI("G")},"Activo",{CLI("I")},">=9")-COUNTIFS({CLI("G")},"Activo",{CLI("I")},"<=6"))/COUNTIFS({CLI("G")},"Activo",{CLI("I")},">=0")*100,0)', 50, "0", "mayor"),
]
FILA = {}
r = 5
for k, *_ in KPI:
    if k: FILA[k] = r
    r += 1
CELDA = {k: f"B{v}" for k, v in FILA.items()}
r = 5
for k, nombre, fx, meta, fmt, regla in KPI:
    if fx is None:
        for j in range(1, 5): t.cell(r, j).fill = PatternFill("solid", fgColor="D9E1F2")
        t.cell(r, 1, nombre).font = Font(bold=True, color="1F3864")
        r += 1; continue
    for kk, cc in CELDA.items():
        fx = fx.replace("{" + kk + "}", cc)
    t.cell(r, 1, nombre)
    c = t.cell(r, 2, fx); c.number_format = fmt
    if meta is not None:
        m = t.cell(r, 3, meta); m.number_format = fmt
    if regla == "mayor":
        t.cell(r, 4, f'=IF(B{r}>=C{r},"Verde",IF(B{r}>=C{r}*0.7,"Amarillo","Rojo"))')
    elif regla == "menor":
        t.cell(r, 4, f'=IF(B{r}<=C{r},"Verde",IF(B{r}<=C{r}*1.5+0.0001,"Amarillo","Rojo"))')
    for j in range(1, 5): t.cell(r, j).border = B
    t.cell(r, 4).alignment = Alignment(horizontal="center")
    r += 1
rng = f"D5:D{r}"
cf(t, rng, "Verde", VERDE); cf(t, rng, "Amarillo", AMAR); cf(t, rng, "Rojo", ROJO)
# Referencia del MRR para la hoja Clientes (% del MRR)
for rr in range(2, N + 2):
    cl[f"L{rr}"] = f'=IF(OR(A{rr}="",G{rr}<>"Activo"),"",IFERROR(E{rr}/Tablero!${CELDA["MRR"][0]}${CELDA["MRR"][1:]},0))'

# Los KPI del tablero deben quedar en filas fijas para las referencias cruzadas
wb.save("plantillas/centro-de-control.xlsx")
