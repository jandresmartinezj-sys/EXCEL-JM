"""Genera la plantilla 'Cartera por edades y deterioro'."""
import datetime as dt
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

AZUL = PatternFill("solid", fgColor="1F3864")
INPUT = Font(color="0000FF")
BLANCO = Font(color="FFFFFF", bold=True)
B = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
COP = '#,##0;[Red]-#,##0'
RANGOS = [("No vencida", 0), ("1-30", 0.01), ("31-60", 0.05), ("61-90", 0.10),
          ("91-180", 0.25), ("181-360", 0.50), ("Más de 360", 1.00)]
N = 500  # filas disponibles en Facturas

wb = Workbook()

# Instrucciones
ws = wb.active; ws.title = "Instrucciones"
txt = [
    ("CARTERA POR EDADES Y DETERIORO", True),
    ("", False),
    ("1. En 'Parametros' define la fecha de corte y los % de deterioro por rango.", False),
    ("2. En 'Facturas' registra cada factura pendiente (columnas en azul). El resto se calcula solo.", False),
    ("3. 'Resumen' muestra la cartera por edades, % de participación, deterioro y alertas.", False),
    ("", False),
    ("IMPORTANTE: los % de deterioro son EJEMPLOS. Ajústalos a la política contable de la empresa:", True),
    ("- NIIF plenas (Grupo 1): modelo de pérdida crediticia esperada (NIIF 9), p. ej. matriz de provisiones con históricos.", False),
    ("- NIIF para las Pymes (Grupo 2): deterioro con evidencia objetiva (Sección 11).", False),
    ("Esta plantilla es una herramienta de apoyo; no reemplaza el juicio profesional del contador.", False),
]
for i, (t, b) in enumerate(txt, 1):
    ws.cell(i, 1, t).font = Font(bold=b, size=14 if i == 1 else 11)
ws.column_dimensions["A"].width = 110

# Parametros
p = wb.create_sheet("Parametros")
p["A1"], p["B1"] = "Fecha de corte", dt.date(2026, 9, 30)
p["B1"].number_format = "dd/mm/yyyy"; p["B1"].font = INPUT; p["A1"].font = Font(bold=True)
p["A3"], p["B3"] = "Rango", "% deterioro"
for c in ("A3", "B3"):
    p[c].fill, p[c].font = AZUL, BLANCO
for i, (r, pct) in enumerate(RANGOS, 4):
    p.cell(i, 1, r).border = B
    c = p.cell(i, 2, pct); c.number_format = "0.0%"; c.font = INPUT; c.border = B
p.column_dimensions["A"].width = 18; p.column_dimensions["B"].width = 14
TAB = f"Parametros!$A$4:$B${3 + len(RANGOS)}"

# Facturas
f = wb.create_sheet("Facturas")
hdr = ["Cliente", "NIT", "N.º factura", "Fecha emisión", "Fecha vencimiento", "Valor",
       "Abonos", "Saldo", "Días mora", "Rango", "% deterioro", "Deterioro"]
for j, h in enumerate(hdr, 1):
    c = f.cell(1, j, h); c.fill, c.font = AZUL, BLANCO
    c.alignment = Alignment(horizontal="center", wrap_text=True)
for col, w in zip("ABCDEFGHIJKL", [32, 14, 12, 13, 13, 15, 13, 15, 10, 12, 11, 15]):
    f.column_dimensions[col].width = w
for r in range(2, N + 2):
    f[f"H{r}"] = f'=IF(F{r}="","",F{r}-N(G{r}))'
    f[f"I{r}"] = f'=IF(OR(F{r}="",E{r}=""),"",MAX(0,Parametros!$B$1-E{r}))'
    f[f"J{r}"] = (f'=IF(OR(H{r}="",I{r}=""),"",IF(H{r}<=0,"Pagada",IF(I{r}=0,"No vencida",'
                  f'IF(I{r}<=30,"1-30",IF(I{r}<=60,"31-60",IF(I{r}<=90,"61-90",'
                  f'IF(I{r}<=180,"91-180",IF(I{r}<=360,"181-360","Más de 360"))))))))')
    f[f"K{r}"] = f'=IF(OR(J{r}="",J{r}="Pagada"),"",VLOOKUP(J{r},{TAB},2,FALSE))'
    f[f"L{r}"] = f'=IF(K{r}="","",H{r}*K{r})'
    for col in "DE": f[f"{col}{r}"].number_format = "dd/mm/yyyy"
    for col in "FGHL": f[f"{col}{r}"].number_format = COP
    f[f"K{r}"].number_format = "0.0%"
    for col in "ABCDEFG": f[f"{col}{r}"].font = INPUT
dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2000,1,1)",
                    error="Ingresa una fecha válida", showErrorMessage=True)
f.add_data_validation(dv); dv.add(f"D2:E{N + 1}")
f.freeze_panes = "A2"; f.auto_filter.ref = f"A1:L{N + 1}"

# Datos de ejemplo (ficticios)
ej = [("Comercializadora Andina SAS", "900111222", 450, 420, 2_400_000, 0),
      ("Distribuciones El Sol SAS", "900333444", 75, 60, 1_850_000, 350_000),
      ("Ferretería La Central", "800555666", 40, 10, 980_000, 0),
      ("Inversiones Río Claro SAS", "901777888", 15, -10, 3_200_000, 0),
      ("Servicios Médicos Norte IPS", "900999000", 200, 170, 5_600_000, 1_000_000),
      ("Tienda Naturista Vida", "1067123456", 420, 400, 760_000, 0),
      ("Comercializadora Andina SAS", "900111222", 30, 5, 1_100_000, 1_100_000),
      ("Constructora Horizonte SAS", "901222333", 300, 270, 8_900_000, 2_000_000)]
for i, (cli, nit, emi, venc, val, ab) in enumerate(ej, 2):
    corte = dt.date(2026, 9, 30)
    f.cell(i, 1, cli); f.cell(i, 2, nit); f.cell(i, 3, f"FE-{1000 + i}")
    f.cell(i, 4, corte - dt.timedelta(days=emi + 30)).number_format = "dd/mm/yyyy"
    f.cell(i, 5, corte - dt.timedelta(days=venc)).number_format = "dd/mm/yyyy"
    f.cell(i, 6, val); f.cell(i, 7, ab or None)

# Resumen
s = wb.create_sheet("Resumen", 1)
s["A1"] = "RESUMEN DE CARTERA"; s["A1"].font = Font(bold=True, size=14)
s["A2"] = "Fecha de corte"; s["B2"] = "=Parametros!B1"; s["B2"].number_format = "dd/mm/yyyy"
for j, h in enumerate(["Rango", "Saldo", "% del total", "N.º facturas", "Deterioro"], 1):
    c = s.cell(4, j, h); c.fill, c.font = AZUL, BLANCO
FR = f"Facturas!$J$2:$J${N + 1}"
for i, (r, _) in enumerate(RANGOS, 5):
    s.cell(i, 1, r)
    s.cell(i, 2, f'=SUMIFS(Facturas!$H$2:$H${N + 1},{FR},A{i})').number_format = COP
    s.cell(i, 3, f'=IF($B$12=0,0,B{i}/$B$12)').number_format = "0.0%"
    s.cell(i, 4, f'=COUNTIFS({FR},A{i})')
    s.cell(i, 5, f'=SUMIFS(Facturas!$L$2:$L${N + 1},{FR},A{i})').number_format = COP
s["A12"] = "TOTAL"; s["A12"].font = Font(bold=True)
for col, fmt in zip("BCDE", [COP, "0.0%", "0", COP]):
    c = s[f"{col}12"]; c.value = f"=SUM({col}5:{col}11)"; c.number_format = fmt; c.font = Font(bold=True)
for r in range(4, 13):
    for col in "ABCDE": s[f"{col}{r}"].border = B

s["A14"] = "INDICADORES"; s["A14"].font = Font(bold=True)
kpi = [("Cartera vencida (%)", "=IF(B12=0,0,1-B5/B12)", "0.0%"),
       ("Cartera > 180 días (%)", "=IF(B12=0,0,(B10+B11)/B12)", "0.0%"),
       ("Deterioro / cartera (%)", "=IF(B12=0,0,E12/B12)", "0.0%"),
       ("Cartera neta de deterioro", "=B12-E12", COP)]
for i, (k, fx, fmt) in enumerate(kpi, 15):
    s.cell(i, 1, k); c = s.cell(i, 2, fx); c.number_format = fmt
s["A20"] = "Alerta"
s["B20"] = ('=IF(B16>0.3,"ALTA: más del 30% de la cartera supera 180 días",'
            'IF(B15>0.5,"MEDIA: más de la mitad de la cartera está vencida","Normal"))')
s["A20"].font = Font(bold=True)
s.column_dimensions["A"].width = 28
for col in "BCDE": s.column_dimensions[col].width = 18

ch = BarChart(); ch.title = "Cartera por edades"; ch.y_axis.title = "COP"
ch.add_data(Reference(s, min_col=2, min_row=4, max_row=11), titles_from_data=True)
ch.set_categories(Reference(s, min_col=1, min_row=5, max_row=11))
ch.legend = None; ch.height, ch.width = 8, 16
s.add_chart(ch, "G4")

wb.save("plantillas/cartera-por-edades-deterioro.xlsx")
